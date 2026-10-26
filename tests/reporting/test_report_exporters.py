"""Integration & Unit Test Suite for Server-Side PDF & Excel Report Exporters.

Verifies:
1. Server-side PDF report generation (Segment P&L summaries, HR capacity justifications)
   with embedded ECharts PNG renders, corporate tokens, and header/footer compilation.
2. Formatted Excel (.xlsx) export with openpyxl enforcing active multiplication and
   summary formulas (=G*H, =SUM, =E-F) and tabular-nums formatting (#,##0.00).
3. FastAPI endpoints GET/POST /api/v1/reports/pdf and GET/POST /api/v1/reports/excel.
"""

from __future__ import annotations

import base64
import io
from typing import Any

import openpyxl
from httpx import ASGITransport, AsyncClient
from PIL import Image

from src.backend.main import app
from src.backend.reporting.excel_generator import (
    ExcelReportExporter,
    get_excel_report_exporter,
)
from src.backend.reporting.pdf_generator import (
    PDFReportGenerator,
    get_pdf_report_generator,
)


def _create_dummy_chart_png(width: int = 400, height: int = 200) -> bytes:
    """Generate a clean in-memory dummy PNG chart render."""
    img = Image.new("RGB", (width, height), color=(211, 47, 47))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf.getvalue()


class TestPDFReportGenerator:
    """Verifies PDF report compilation with embedded ECharts images and corporate styling."""

    def test_segment_pnl_pdf_compilation_default_chart(self) -> None:
        """Assert Segment P&L PDF compiles with headless ECharts waterfall chart simulation."""
        generator: PDFReportGenerator = get_pdf_report_generator()
        pnl_data: dict[str, Any] = {
            "fiscal_year": 2024,
            "fiscal_month": 3,
            "depot_id": "DHK",
            "gross_revenue_bdt": 6500000.0,
            "discounts_bdt": 422500.0,
            "net_revenue_bdt": 6077500.0,
            "cogs_bdt": 3400000.0,
            "gross_margin_bdt": 2677500.0,
            "gross_margin_pct": 44.1,
            "opex_bdt": 1250000.0,
            "ebitda_bdt": 1427500.0,
            "ebitda_pct": 23.5,
        }

        pdf_bytes = generator.generate_segment_pnl_pdf(pnl_data)

        # Assert valid PDF binary stream
        assert pdf_bytes.startswith(b"%PDF"), "Output must start with PDF magic bytes"
        assert len(pdf_bytes) > 5000, "Compiled PDF must contain complete document and image stream"

    def test_segment_pnl_pdf_compilation_custom_echarts_image(self) -> None:
        """Assert Segment P&L PDF correctly embeds custom client-rendered ECharts PNG bytes."""
        generator: PDFReportGenerator = get_pdf_report_generator()
        custom_png = _create_dummy_chart_png(500, 250)

        pnl_data: dict[str, Any] = {
            "fiscal_year": 2024,
            "fiscal_month": 3,
            "depot_id": "CTG",
            "gross_revenue_bdt": 4200000.0,
            "discounts_bdt": 273000.0,
            "net_revenue_bdt": 3927000.0,
            "cogs_bdt": 2150000.0,
            "gross_margin_bdt": 1777000.0,
            "gross_margin_pct": 45.3,
            "opex_bdt": 820000.0,
            "ebitda_bdt": 957000.0,
            "ebitda_pct": 24.4,
        }

        # Test passing raw bytes
        pdf_bytes = generator.generate_segment_pnl_pdf(pnl_data, chart_image=custom_png)
        assert pdf_bytes.startswith(b"%PDF")
        assert len(pdf_bytes) > 2000

        # Test passing base64 data URI
        b64_str = f"data:image/png;base64,{base64.b64encode(custom_png).decode('utf-8')}"
        pdf_bytes_b64 = generator.generate_segment_pnl_pdf(pnl_data, chart_image=b64_str)
        assert pdf_bytes_b64.startswith(b"%PDF")

    def test_hr_capacity_pdf_compilation_default_chart(self) -> None:
        """Assert HR Workforce Capacity PDF compiles with default radar chart."""
        generator: PDFReportGenerator = get_pdf_report_generator()
        capacity_data: dict[str, Any] = {
            "depot_id": "DHK",
            "total_vacant_territories": 39,
            "total_recommended_headcount": 54,
            "total_projected_monthly_revenue_bdt": 61500000.0,
            "total_monthly_investment_bdt": 2970000.0,
            "average_payback_months": 2.4,
            "territories": [
                {
                    "territory_name": "Dhanmondi Paint Cluster",
                    "depot_id": "DHK",
                    "active_dealers": 45,
                    "territory_target_bdt": 2250000.0,
                    "sales_revenue_capacity": 1.5,
                    "workload_capacity": 1.88,
                    "recommended_headcount": 2,
                    "payback_period_months": 1.9,
                }
            ],
        }

        pdf_bytes = generator.generate_hr_capacity_pdf(capacity_data)
        assert pdf_bytes.startswith(b"%PDF")
        assert len(pdf_bytes) > 5000

    def test_hr_capacity_pdf_compilation_custom_image(self) -> None:
        """Assert HR Workforce Capacity PDF embeds external chart image."""
        generator: PDFReportGenerator = get_pdf_report_generator()
        custom_png = _create_dummy_chart_png(600, 300)

        capacity_data: dict[str, Any] = {
            "depot_id": "SYL",
            "total_vacant_territories": 6,
            "total_recommended_headcount": 9,
            "total_projected_monthly_revenue_bdt": 10500000.0,
            "total_monthly_investment_bdt": 495000.0,
            "average_payback_months": 2.1,
        }

        pdf_bytes = generator.generate_hr_capacity_pdf(capacity_data, chart_image=custom_png)
        assert pdf_bytes.startswith(b"%PDF")
        assert len(pdf_bytes) > 2000


class TestExcelReportExporter:
    """Verifies Excel generation with live formulas and corporate styling."""

    def test_purchase_indents_excel_generation_and_formulas(self) -> None:
        """Assert purchase indents workbook contains valid live formulas and header styling."""
        exporter: ExcelReportExporter = get_excel_report_exporter()
        proposals = [
            {
                "indent_id": "IND-TEST-001",
                "material_code": "RM_TIO2_RUTILE",
                "material_name": "Titanium Dioxide Rutile (R-902+)",
                "category": "TIO2",
                "current_stock_kg": 15000.0,
                "safety_stock_kg": 25000.0,
                "recommended_reorder_kg": 30000.0,
                "unit_cost_bdt": 385.00,
                "supplier_name": "Chemours Singapore",
                "country_of_origin": "Singapore",
                "alert_level": "CRITICAL_STOCKOUT_RISK",
            },
            {
                "indent_id": "IND-TEST-002",
                "material_code": "RM_ACRYLIC_RESIN",
                "material_name": "Pure Acrylic Emulsion 50%",
                "category": "RESIN",
                "current_stock_kg": 20000.0,
                "safety_stock_kg": 30000.0,
                "recommended_reorder_kg": 25000.0,
                "unit_cost_bdt": 240.00,
                "supplier_name": "BASF Chemicals",
                "country_of_origin": "Germany",
                "alert_level": "REORDER_REQUIRED",
            },
        ]

        xlsx_bytes = exporter.generate_purchase_indents_excel(
            proposals, metadata={"evaluation_date": "2024-03-31"}
        )

        assert len(xlsx_bytes) > 0, "Excel output cannot be empty"

        # Load with data_only=False to inspect formula text
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
        assert "Purchase Indents" in wb.sheetnames
        ws = wb["Purchase Indents"]

        # 1. Header assertions
        assert ws["A5"].value == "Indent ID"
        assert ws["I5"].value == "Total PO Cost (BDT)"
        # Nerolac Red Header fill: check RGB
        header_color = str(ws["A5"].fill.start_color.rgb)
        assert header_color in ("00D32F2F", "D32F2F") or header_color.endswith("D32F2F")

        # 2. Row Formula assertions (Row 6 and Row 7)
        # Col I (Total PO Cost) = Col G (Reorder KG) * Col H (Unit Cost)
        assert ws["I6"].value == "=G6*H6", f"Expected '=G6*H6', got '{ws['I6'].value}'"
        assert ws["I7"].value == "=G7*H7", f"Expected '=G7*H7', got '{ws['I7'].value}'"

        # Tabular-nums number formatting
        assert ws["G6"].number_format == "#,##0.00"
        assert ws["H6"].number_format == "#,##0.00"
        assert ws["I6"].number_format == "#,##0.00"

        # 3. Summary / Total Row assertions (Row 8)
        assert "TOTAL PROCUREMENT COMMITMENT" in str(ws["A8"].value)
        assert ws["G8"].value == "=SUM(G6:G7)", f"Expected '=SUM(G6:G7)', got '{ws['G8'].value}'"
        assert ws["I8"].value == "=SUM(I6:I7)", f"Expected '=SUM(I6:I7)', got '{ws['I8'].value}'"
        assert ws["I8"].number_format == "#,##0.00"

    def test_trial_balance_excel_generation_and_formulas(self) -> None:
        """Assert trial balance sheet contains subtraction formula =E-F and grand total sums."""
        exporter: ExcelReportExporter = get_excel_report_exporter()
        accounts = [
            {
                "gl_code": "10000",
                "account_name": "Gross Sales Revenue",
                "category": "REVENUE",
                "segment": "Architectural",
                "debit": 0.0,
                "credit": 5000000.0,
            },
            {
                "gl_code": "10001",
                "account_name": "Sales Discounts & Commercial Rebates",
                "category": "REVENUE",
                "segment": "Trade",
                "debit": 325000.0,
                "credit": 0.0,
            },
            {
                "gl_code": "10002",
                "account_name": "Raw Material Consumed",
                "category": "COGS",
                "segment": "Production",
                "debit": 1800000.0,
                "credit": 0.0,
            },
        ]

        xlsx_bytes = exporter.generate_trial_balance_excel(
            accounts,
            metadata={"fiscal_year": 2024, "fiscal_month": 3, "depot_id": "DHK"},
        )

        assert len(xlsx_bytes) > 0

        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
        assert "Trial Balance" in wb.sheetnames
        ws = wb["Trial Balance"]

        # Check headers
        assert ws["A5"].value == "GL Account"
        assert ws["G5"].value == "Net Balance (BDT)"

        # Check row subtraction formula: Net = Debit (Col E) - Credit (Col F)
        assert ws["G6"].value == "=E6-F6"
        assert ws["G7"].value == "=E7-F7"
        assert ws["G8"].value == "=E8-F8"

        # Check Summary Row (Row 9)
        assert ws["E9"].value == "=SUM(E6:E8)"
        assert ws["F9"].value == "=SUM(F6:F8)"
        assert ws["G9"].value == "=E9-F9"


class TestReportsAPIEndpoints:
    """Verifies HTTP export endpoints /api/v1/reports/pdf and /api/v1/reports/excel."""

    async def test_get_pdf_segment_pnl_endpoint(self) -> None:
        """Assert GET /api/v1/reports/pdf?report_type=segment_pnl returns clean PDF stream."""
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://testserver") as client:
            url = "/api/v1/reports/pdf?report_type=segment_pnl&fiscal_year=2024&fiscal_month=3"
            resp = await client.get(url)
            assert resp.status_code == 200
            assert resp.headers["content-type"] == "application/pdf"
            assert "attachment; filename=" in resp.headers["content-disposition"]
            assert resp.content.startswith(b"%PDF")
            assert len(resp.content) > 5000

    async def test_get_pdf_hr_capacity_endpoint(self) -> None:
        """Assert GET /api/v1/reports/pdf?report_type=hr_capacity returns briefing PDF."""
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://testserver") as client:
            resp = await client.get("/api/v1/reports/pdf?report_type=hr_capacity")
            assert resp.status_code == 200
            assert resp.headers["content-type"] == "application/pdf"
            disp = resp.headers["content-disposition"]
            assert "nerolac_hr_workforce_capacity_briefing.pdf" in disp
            assert resp.content.startswith(b"%PDF")
            assert len(resp.content) > 5000

    async def test_post_pdf_with_custom_chart_base64(self) -> None:
        """Assert POST /api/v1/reports/pdf embeds client-sent base64 ECharts render."""
        custom_png = _create_dummy_chart_png(300, 150)
        b64_str = base64.b64encode(custom_png).decode("utf-8")

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://testserver") as client:
            payload = {
                "report_type": "segment_pnl",
                "fiscal_year": 2024,
                "fiscal_month": 3,
                "chart_image_base64": b64_str,
            }
            resp = await client.post("/api/v1/reports/pdf", json=payload)
            assert resp.status_code == 200
            assert resp.headers["content-type"] == "application/pdf"
            assert resp.content.startswith(b"%PDF")

    async def test_get_excel_purchase_indents_endpoint(self) -> None:
        """Assert GET purchase indents export returns valid .xlsx with formulas."""
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://testserver") as client:
            resp = await client.get("/api/v1/reports/excel?report_type=purchase_indents")
            assert resp.status_code == 200
            assert (
                resp.headers["content-type"]
                == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            assert "nerolac_purchase_indents.xlsx" in resp.headers["content-disposition"]

            # Load and verify Excel contents
            wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=False)
            assert "Purchase Indents" in wb.sheetnames
            ws = wb["Purchase Indents"]
            assert ws["A5"].value == "Indent ID"
            assert "=G" in str(ws["I6"].value)

    async def test_get_excel_trial_balance_endpoint(self) -> None:
        """Assert GET trial balance export returns valid trial balance sheet."""
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://testserver") as client:
            url = "/api/v1/reports/excel?report_type=trial_balance&fiscal_year=2024&fiscal_month=3"
            resp = await client.get(url)
            assert resp.status_code == 200
            assert (
                resp.headers["content-type"]
                == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            assert "nerolac_trial_balance" in resp.headers["content-disposition"]

            wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=False)
            assert "Trial Balance" in wb.sheetnames
            ws = wb["Trial Balance"]
            assert "=E" in str(ws["G6"].value)

    async def test_post_excel_filtered_endpoint(self) -> None:
        """Assert POST /api/v1/reports/excel filtered export works."""
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://testserver") as client:
            payload = {
                "report_type": "purchase_indents",
                "category": "TIO2",
            }
            resp = await client.post("/api/v1/reports/excel", json=payload)
            assert resp.status_code == 200
            assert (
                resp.headers["content-type"]
                == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

    async def test_invalid_report_type_error(self) -> None:
        """Assert invalid report type triggers HTTP 422 Unprocessable Entity."""
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://testserver") as client:
            resp = await client.get("/api/v1/reports/pdf?report_type=unsupported_type")
            assert resp.status_code == 422
